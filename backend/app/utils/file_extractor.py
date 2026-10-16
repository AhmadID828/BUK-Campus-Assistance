from pypdf import PdfReader
from docx import Document
from pptx import Presentation #type: ignore
from sqlalchemy import text
from typing import Optional
import os

def extract_text_from_file(file_path: str) -> str:
    """Extract text from PDF, DOCX or PPTX files"""
    extension = os.path.splitext(file_path)[1].lower()

    try:
        if extension == ".pdf":
            reader = PdfReader(file_path)
            text = ""
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
            return text.strip()

        elif extension == ".docx":
            doc = Document(file_path)
            text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
            return text.strip()

        elif extension == ".pptx":
            prs = Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip(): #type: ignore
                        text += shape.text + "\n" #type: ignore
            return text.strip()

        else:
            return "Error extracting text: Unsupported file type"

    except Exception as e:
        return f"Error extracting text: {str(e)}"
    
def update_search_vector(db, material_id: Optional[int] = None) -> None:
    """
    Update search_vector for one material or all materials
    """
    if material_id is not None:
        # Update single material
        db.execute(
            text("""
                UPDATE study_materials 
                SET search_vector = 
                    setweight(to_tsvector('english', coalesce(title, '')), 'A') ||
                    setweight(to_tsvector('english', coalesce(subject, '')), 'B') ||
                    setweight(to_tsvector('english', coalesce(content, '')), 'C')
                WHERE id = :id
            """),
            {"id": material_id}
        )
    else:
        # Update ALL materials
        db.execute(
            text("""
                UPDATE study_materials 
                SET search_vector = 
                    setweight(to_tsvector('english', coalesce(title, '')), 'A') ||
                    setweight(to_tsvector('english', coalesce(subject, '')), 'B') ||
                    setweight(to_tsvector('english', coalesce(content, '')), 'C')
            """)
        )
    db.commit()
